"""Evaluate saved formulas independently of their cached values; never modify the delivered file."""
import json, math, statistics, sys
from collections import Counter
from pathlib import Path
import openpyxl
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from formula_engine import Engine, Error, isnum

class ParentEngine(Engine):
    def eval(self,n,s):
        if n[0]=='fn':
            name,args=n[1:]
            if name in ['COUNT','MEDIAN','MODE']:
                vals=[v for arg in args for v in self.values(arg,s) if isnum(v)]
                if name=='COUNT':return len(vals)
                if not vals:raise Error('EMPTY '+name)
                if name=='MEDIAN':return statistics.median(vals)
                counts=Counter(vals);mx=max(counts.values())
                if mx==1:raise Error('N/A MODE')
                return next(v for v in vals if counts[v]==mx)
            if name=='IFERROR':
                try:return self.eval(args[0],s)
                except Error as err:
                    if str(err)!='DIV/0':raise
                    return self.eval(args[1],s)
        return super().eval(n,s)

root=Path(__file__).resolve().parents[1]
out=root/'outputs'/'20260920-new-source-parent-model'
book=out/'户外地垫市场分析-SPEC3-父体口径可回勾.xlsx'
wf=openpyxl.load_workbook(book,read_only=True,data_only=False)
wc=openpyxl.load_workbook(book,read_only=True,data_only=True)
sheets={}; formulas=[]
for ws in wf:
    if ws.title in ['90_原始输入','93_来源登记','94_规则说明','00_数据总览']:continue
    cells={};sheets[ws.title]=cells
    for row,cached in zip(ws.iter_rows(),wc[ws.title].iter_rows()):
        for cell,cc in zip(row,cached):
            if cell.value is None:continue
            key=(cell.row,cell.column)
            if cell.data_type=='f':
                cells[key]={'f':cell.value};formulas.append((ws.title,*key,cc.value))
            else:cells[key]=cell.value
engine=ParentEngine(sheets)
def equal(a,b):
    if a in [None,''] and b in [None,'']:return True
    if isnum(a) and isnum(b):return math.isclose(a,b,abs_tol=1e-7,rel_tol=1e-10)
    return a==b
failures=[]
for sheet,r,c,cached in formulas:
    actual=engine.get(sheet,r,c)
    if not equal(actual,cached):
        failures.append({'sheet':sheet,'row':r,'col':c,'formula':sheets[sheet][(r,c)]['f'],'actual':actual,'cached':cached})
        if len(failures)>=12:break
print(json.dumps({'formulas':len(formulas),'evaluated':engine.count,'failures':failures},ensure_ascii=False),flush=True)
assert not failures

# Perturb a single valid child amount in an in-memory copy. Aggregate revenue and ASP must change.
model=json.loads((out/'model.json').read_text(encoding='utf-8'))
sample=next(p for p in model['parentRows'] if p['month']=='202607' and p['parentKey']=='B0BM74444Q')
pi=model['parentRows'].index(sample)+2
ci=next(i+2 for i,r in enumerate(model['candidateRows']) if r['month']=='202607' and r['parentKey']=='B0BM74444Q' and r['childSales'] is not None and r['childRevenue'] is not None)
detail=sheets['91_BSR候选子体明细'];old=detail[(ci,12)];detail[(ci,12)]=old+100
changed=ParentEngine(sheets)
assert equal(changed.get('92_父体月度汇总',pi,20),sample['childRevenue']+100)
assert equal(changed.get('92_父体月度汇总',pi,23),(sample['childRevenue']+100)/sample['childSales'])
assert equal(changed.get('01_整体市场月度',25,6),model['summaries']['overall'][-1]['childRevenue']+100)
detail[(ci,12)]=old

# Remove both fields from all children of the sample: empty aggregate, not zero/-100%.
sample_indices=[i+2 for i,r in enumerate(model['candidateRows']) if r['month']=='202607' and r['parentKey']=='B0BM74444Q']
for row in sample_indices: detail[(row,11)]=None; detail[(row,12)]=None
changed=ParentEngine(sheets)
assert changed.get('92_父体月度汇总',pi,20)==''
assert changed.get('92_父体月度汇总',pi,22)=='NONE'
# Explicit zero remains numeric zero.
detail[(sample_indices[0],11)]=0; detail[(sample_indices[0],12)]=0
changed=ParentEngine(sheets)
assert changed.get('92_父体月度汇总',pi,20)==0
assert changed.get('92_父体月度汇总',pi,18)==1
wf.close();wc.close()
result={'status':'PASS','savedFormulasEvaluated':len(formulas),'inputChanges':['child revenue +100','all child metrics blank','explicit zero'],'engine':'bounded independent evaluator; not desktop Excel'}
(root/'tmp'/'parent-review'/'formula-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(result,ensure_ascii=False))
